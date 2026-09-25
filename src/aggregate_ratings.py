


import os
import re
import sys
import glob
import argparse

import numpy as np
import pandas as pd
from openpyxl import load_workbook

try:
    from scipy import stats as scipy_stats
    _HAS_SCIPY = True
except ImportError:
    _HAS_SCIPY = False

LEVEL1_CATEGORIES = ["supported", "partial", "unsupported", "unknown"]
LEVEL2_DIMS = ["faithfulness", "actionability", "clarity"]
NEGATIVE_CONTROL_GATE = 0.70  



TAG_NORMALIZE = {
    '지지됨': 'supported', '부분지지': 'partial', '근거없음': 'unsupported', '판단불가': 'unknown',
    'supported': 'supported', 'partial': 'partial', 'unsupported': 'unsupported', 'unknown': 'unknown',
}


# ════════════════════════════════════════════════════════════════

# ════════════════════════════════════════════════════════════════

def _normalize_tag(raw):
    if raw is None:
        return None
    raw = str(raw).strip()
    return TAG_NORMALIZE.get(raw, raw if raw else None)


def _parse_level1_tags_field(tag_str: str) -> dict:


    result = {}
    if not isinstance(tag_str, str) or not tag_str.strip():
        return result
    for part in tag_str.split(';'):
        part = part.strip()
        if not part or ':' not in part:
            continue
        sid, tag = part.split(':', 1)
        result[sid.strip()] = _normalize_tag(tag)
    return result


def load_csv_response(path: str):
    df = pd.read_csv(path, dtype=str)
    required = {'item_id', 'rater_id', 'level1_tags', 'faithfulness', 'actionability', 'clarity'}
    missing = required - set(df.columns)
    if missing:
        print(f"[경고] {path}: 예상 컬럼이 없습니다({missing}) — rater_tool.html CSV 다운로드 산출물이 "
              f"맞는지 확인하세요. 이 파일은 건너뜁니다.")
        return None

    rows = []
    for _, r in df.iterrows():
        tags = _parse_level1_tags_field(r.get('level1_tags', ''))
        base = {
            'item_id': r['item_id'], 'rater_id': r['rater_id'],
            'faithfulness': r.get('faithfulness'), 'actionability': r.get('actionability'),
            'clarity': r.get('clarity'), 'free_text': r.get('free_text', ''),
            'duration_sec': r.get('duration_sec'), '__source_file': os.path.basename(path),
        }
        if tags:
            for sid, tag in tags.items():
                rows.append({**base, 'sentence_id': sid, 'level1_tag': tag})
        else:
            
            rows.append({**base, 'sentence_id': None, 'level1_tag': None})
    return pd.DataFrame(rows)


def load_xlsx_response(path: str):
    wb = load_workbook(path, data_only=True)
    if 'Level1_문장평정' not in wb.sheetnames or 'Level2_문항평정' not in wb.sheetnames:
        print(f"[경고] {path}: 'Level1_문장평정'/'Level2_문항평정' 시트가 없습니다 — "
              f"export_rater_form_to_excel.py 산출물이 맞는지 확인하세요. 건너뜁니다.")
        return None

    fname = os.path.basename(path)
    m = re.match(r'([^_]+)_평정양식', fname)
    rater_id = m.group(1) if m else None
    if not rater_id and '안내' in wb.sheetnames:
        for row in wb['안내'].iter_rows(min_row=1, max_row=10, max_col=1):
            v = row[0].value
            if isinstance(v, str) and '평정자 ID' in v:
                rater_id = v.split(':', 1)[-1].strip()
                break
    if not rater_id:
        print(f"[경고] {path}: 평정자 ID를 파일명이나 '안내' 시트에서 찾지 못했습니다 — 건너뜁니다. "
              f"파일명을 '<rater_id>_평정양식.xlsx' 형태로 유지해 회신받는 것을 권장합니다.")
        return None

    ws1 = wb['Level1_문장평정']
    headers1 = [c.value for c in next(ws1.iter_rows(min_row=1, max_row=1))]
    rows_l1 = []
    for row in ws1.iter_rows(min_row=2, values_only=True):
        d = dict(zip(headers1, row))
        item_id = d.get('문항ID')
        if not item_id or str(item_id).startswith('(예시)'):
            continue
        tag_raw = d.get('평정(드롭다운 선택)')
        if tag_raw is None or str(tag_raw).strip() == '':
            continue
        rows_l1.append({'item_id': item_id, 'sentence_id': d.get('문장ID'),
                         'level1_tag': _normalize_tag(tag_raw)})
    l1_df = pd.DataFrame(rows_l1)

    ws2 = wb['Level2_문항평정']
    headers2 = [c.value for c in next(ws2.iter_rows(min_row=1, max_row=1))]
    rows_l2 = []
    for row in ws2.iter_rows(min_row=2, values_only=True):
        d = dict(zip(headers2, row))
        item_id = d.get('문항ID')
        if not item_id or str(item_id).startswith('(예시)'):
            continue
        rows_l2.append({
            'item_id': item_id,
            'faithfulness': d.get('충실도(1-5)'), 'actionability': d.get('실행가능성(1-5)'),
            'clarity': d.get('명료성(1-5)'), 'free_text': d.get('자유서술(선택)') or '',
        })
    l2_df = pd.DataFrame(rows_l2)

    if l1_df.empty and l2_df.empty:
        print(f"[경고] {path}: 채워진 응답이 없습니다 — 건너뜁니다.")
        return None
    if l1_df.empty:
        merged = l2_df.copy()
        merged['sentence_id'] = None
        merged['level1_tag'] = None
    elif l2_df.empty:
        merged = l1_df.copy()
        for c in LEVEL2_DIMS:
            merged[c] = None
        merged['free_text'] = ''
    else:
        merged = l1_df.merge(l2_df, on='item_id', how='outer')

    merged['rater_id'] = rater_id
    merged['duration_sec'] = None
    merged['__source_file'] = fname
    return merged


def load_responses(responses_dir: str) -> pd.DataFrame:
    csv_paths = sorted(glob.glob(os.path.join(responses_dir, "*.csv")))
    xlsx_paths = sorted(glob.glob(os.path.join(responses_dir, "*.xlsx")))
    if not csv_paths and not xlsx_paths:
        raise FileNotFoundError(f"'{responses_dir}'에서 .csv/.xlsx 응답 파일을 찾지 못했습니다.")

    frames = []
    for p in csv_paths:
        df = load_csv_response(p)
        if df is not None and not df.empty:
            frames.append(df)
    for p in xlsx_paths:
        df = load_xlsx_response(p)
        if df is not None and not df.empty:
            frames.append(df)
    if not frames:
        raise ValueError("유효한 응답 파일을 하나도 읽지 못했습니다.")

    out = pd.concat(frames, ignore_index=True)
    for col in LEVEL2_DIMS:
        out[col] = pd.to_numeric(out[col], errors="coerce")
    n_raters = out["rater_id"].nunique()
    print(f"[로딩] 응답 파일 {len(frames)}개(csv {len(csv_paths)}개 + xlsx {len(xlsx_paths)}개 중 유효분), "
          f"평정자 {n_raters}명({sorted(out['rater_id'].unique())}), 총 {len(out)}행")
    return out


def load_answer_key(path: str) -> pd.DataFrame:
    key_df = pd.read_csv(path, dtype={"display_id": str})
    key_df["is_negative_control"] = key_df["is_negative_control"].astype(str).str.lower() == "true"
    key_df["in_core_set"] = key_df["in_core_set"].astype(str).str.lower() == "true"
    return key_df.set_index("display_id")


def load_sentence_answer_key(path: str) -> pd.DataFrame:
    df = pd.read_csv(path, dtype={"display_id": str, "sentence_id": str})
    return df


# ════════════════════════════════════════════════════════════════

# ════════════════════════════════════════════════════════════════

def fleiss_kappa(subject_category_counts: np.ndarray) -> dict:


    N, k = subject_category_counts.shape
    n = subject_category_counts.sum(axis=1)
    if not np.all(n == n[0]):
        raise ValueError("모든 subject가 동일한 평정 수를 가져야 Fleiss' kappa를 계산할 수 있습니다.")
    n = int(n[0])
    if n < 2:
        raise ValueError("Fleiss' kappa는 subject당 평정자 수가 2 이상이어야 합니다.")

    p_j = subject_category_counts.sum(axis=0) / (N * n)
    P_i = (np.sum(subject_category_counts ** 2, axis=1) - n) / (n * (n - 1))
    P_bar = P_i.mean()
    P_e = np.sum(p_j ** 2)
    kappa = (P_bar - P_e) / (1 - P_e) if (1 - P_e) != 0 else float("nan")
    return {"kappa": kappa, "n_subjects": N, "n_raters_per_subject": n, "P_bar": P_bar, "P_e": P_e}


_ORDINAL_VALUE = {'unsupported': 0, 'partial': 1, 'supported': 2}  


def _ordinal_distance(a: str, b: str) -> float:
    if a == b:
        return 0.0
    if a == 'unknown' or b == 'unknown':
        return 2.0  
    return abs(_ORDINAL_VALUE[a] - _ORDINAL_VALUE[b])


def weighted_pairwise_kappa(wide: pd.DataFrame) -> dict:



    raters = list(wide.columns)
    max_dist = 2.0
    pair_kappas = []
    for i in range(len(raters)):
        for j in range(i + 1, len(raters)):
            pair = wide[[raters[i], raters[j]]].dropna()
            if len(pair) < 2:
                continue
            a_vals, b_vals = pair.iloc[:, 0].tolist(), pair.iloc[:, 1].tolist()
            categories = sorted(set(a_vals) | set(b_vals))
            if len(categories) < 2:
                continue
            n = len(a_vals)
            
            obs_weights = [1 - (_ordinal_distance(a, b) / max_dist) ** 2 for a, b in zip(a_vals, b_vals)]
            po_w = np.mean(obs_weights)
            
            pa = {c: a_vals.count(c) / n for c in categories}
            pb = {c: b_vals.count(c) / n for c in categories}
            pe_w = sum(pa[c1] * pb[c2] * (1 - (_ordinal_distance(c1, c2) / max_dist) ** 2)
                       for c1 in categories for c2 in categories)
            if pe_w == 1:
                continue
            pair_kappas.append((po_w - pe_w) / (1 - pe_w))
    if not pair_kappas:
        return {'weighted_kappa': None, 'n_rater_pairs': 0}
    return {'weighted_kappa': float(np.mean(pair_kappas)), 'n_rater_pairs': len(pair_kappas)}


def icc_2_1(ratings: np.ndarray) -> dict:

    n, k = ratings.shape
    if n < 2 or k < 2:
        raise ValueError("ICC(2,1) 계산에는 대상(item) 2개 이상, 평정자 2명 이상이 필요합니다.")

    mean_targets = ratings.mean(axis=1)
    mean_raters = ratings.mean(axis=0)
    grand_mean = ratings.mean()

    SSR = k * np.sum((mean_targets - grand_mean) ** 2)
    SSC = n * np.sum((mean_raters - grand_mean) ** 2)
    SST = np.sum((ratings - grand_mean) ** 2)
    SSE = SST - SSR - SSC

    MSR = SSR / (n - 1)
    MSC = SSC / (k - 1)
    MSE = SSE / ((n - 1) * (k - 1))

    denom = MSR + (k - 1) * MSE + (k / n) * (MSC - MSE)
    icc = (MSR - MSE) / denom if denom != 0 else float("nan")
    return {"icc": icc, "n_targets": n, "k_raters": k, "MSR": MSR, "MSC": MSC, "MSE": MSE}


def pearson_spearman(x: np.ndarray, y: np.ndarray) -> dict:
    mask = ~(np.isnan(x) | np.isnan(y))
    x, y = x[mask], y[mask]
    result = {"n": int(mask.sum())}
    if result["n"] < 3:
        result.update({"pearson_r": float("nan"), "pearson_p": float("nan"),
                        "spearman_rho": float("nan"), "spearman_p": float("nan")})
        return result
    if _HAS_SCIPY:
        r, p = scipy_stats.pearsonr(x, y)
        rho, sp = scipy_stats.spearmanr(x, y)
    else:
        r = float(np.corrcoef(x, y)[0, 1])
        p = float("nan")
        rank_x, rank_y = pd.Series(x).rank(), pd.Series(y).rank()
        rho = float(np.corrcoef(rank_x, rank_y)[0, 1])
        sp = float("nan")
    result.update({"pearson_r": r, "pearson_p": p, "spearman_rho": rho, "spearman_p": sp})
    return result


# ════════════════════════════════════════════════════════════════

# ════════════════════════════════════════════════════════════════

def build_item_level_summary(resp_df: pd.DataFrame, key_df: pd.DataFrame) -> pd.DataFrame:

    level2_first_rows = resp_df.groupby(["item_id", "rater_id"], as_index=False).first()
    agg = level2_first_rows.groupby("item_id").agg(
        n_raters=("rater_id", "nunique"),
        mean_faithfulness=("faithfulness", "mean"),
        mean_actionability=("actionability", "mean"),
        mean_clarity=("clarity", "mean"),
    ).reset_index()

    joined = agg.join(key_df, on="item_id", how="left")
    unmatched = joined[joined["category"].isna()]["item_id"].tolist()
    if unmatched:
        print(f"[경고] 정답키(internal_answer_key.csv)에 없는 item_id {len(unmatched)}개 발견 "
              f"(오타이거나 잘못된 배치 파일일 수 있음): {unmatched[:10]}"
              + (" ..." if len(unmatched) > 10 else ""))

    no_response_ids = set(key_df.index) - set(agg["item_id"])
    if no_response_ids:
        print(f"[경고] answer-key에는 있지만 응답이 하나도 없는 문항 {len(no_response_ids)}개 "
              f"({len(key_df)}개 중) — 특정 평정자가 응답을 안 보냈거나 아직 안 끝냈을 수 있습니다. "
              f"예: {sorted(no_response_ids)[:10]}" + (" ..." if len(no_response_ids) > 10 else ""))
    return joined


def build_sentence_level_summary(resp_df: pd.DataFrame, sentence_key_df: pd.DataFrame) -> pd.DataFrame:


    l1 = resp_df.dropna(subset=["sentence_id", "level1_tag"])
    if l1.empty:
        return pd.DataFrame()

    def _mode_or_none(s):
        m = s.mode()
        return m.iloc[0] if not m.empty else None

    def _counts_str(s):
        return ';'.join(f"{k}={v}" for k, v in s.value_counts().items())

    agg = l1.groupby(["item_id", "sentence_id"]).agg(
        n_raters=("rater_id", "nunique"),
        majority_tag=("level1_tag", _mode_or_none),
        tag_distribution=("level1_tag", _counts_str),
    ).reset_index()

    if sentence_key_df is not None and not sentence_key_df.empty:
        key_small = sentence_key_df.rename(columns={"display_id": "item_id"})[
            ["item_id", "sentence_id", "sentence_en", "auto_status", "auto_entailment", "cited_indicators"]
        ]
        agg = agg.merge(key_small, on=["item_id", "sentence_id"], how="left")
    return agg


def compute_core_set_reliability(resp_df: pd.DataFrame, key_df: pd.DataFrame) -> dict:
    core_ids = set(key_df[key_df["in_core_set"]].index)
    core_resp = resp_df[resp_df["item_id"].isin(core_ids)]
    if core_resp.empty:
        return {"error": "core set 응답이 없습니다 (아직 회신이 오지 않았거나 item_id가 일치하지 않습니다)."}

    result = {}

    
    l1 = core_resp.dropna(subset=["level1_tag"])
    counts = l1.groupby(["item_id", "sentence_id", "level1_tag"]).size().unstack(fill_value=0)
    for cat in LEVEL1_CATEGORIES:
        if cat not in counts.columns:
            counts[cat] = 0
    counts = counts[LEVEL1_CATEGORIES]
    n_per_subject = counts.sum(axis=1)
    modal_n = n_per_subject.mode().iloc[0] if not n_per_subject.empty else 0
    complete = counts[n_per_subject == modal_n]
    dropped = len(counts) - len(complete)
    if len(complete) >= 2 and modal_n >= 2:
        fk = fleiss_kappa(complete.to_numpy())
        fk["n_subjects_dropped_incomplete"] = int(dropped)
        result["level1_fleiss_kappa"] = fk
    else:
        result["level1_fleiss_kappa"] = {"error": "완전한(전원 응답) core set 문장이 충분하지 않습니다."}

    
    l1_wide = l1.pivot_table(index=["item_id", "sentence_id"], columns="rater_id",
                              values="level1_tag", aggfunc="first")
    wk = weighted_pairwise_kappa(l1_wide)
    result["level1_weighted_kappa"] = wk

    # --- Level 2: ICC(2,1) ---
    l2 = core_resp.groupby(["item_id", "rater_id"], as_index=False).first()
    icc_results = {}
    for dim in LEVEL2_DIMS:
        pivot = l2.pivot(index="item_id", columns="rater_id", values=dim)
        complete_rows = pivot.dropna(axis=0, how="any")
        if complete_rows.shape[0] >= 2 and complete_rows.shape[1] >= 2:
            icc_results[dim] = icc_2_1(complete_rows.to_numpy(dtype=float))
            icc_results[dim]["n_items_dropped_incomplete"] = int(pivot.shape[0] - complete_rows.shape[0])
        else:
            icc_results[dim] = {"error": "완전한(전원 응답) core set 문항이 충분하지 않습니다."}
    result["level2_icc"] = icc_results

    return result


def compute_convergent_validity(item_summary: pd.DataFrame) -> dict:
    valid = item_summary.dropna(subset=["mean_faithfulness", "mean_entailment", "faithfulness_rate"])
    return {
        "vs_mean_entailment": pearson_spearman(
            valid["mean_faithfulness"].to_numpy(dtype=float),
            valid["mean_entailment"].to_numpy(dtype=float)),
        "vs_auto_faithfulness_rate": pearson_spearman(
            valid["mean_faithfulness"].to_numpy(dtype=float),
            valid["faithfulness_rate"].to_numpy(dtype=float)),
    }


def compute_sentence_level_crosstab(sentence_summary: pd.DataFrame):


    if sentence_summary is None or sentence_summary.empty or "auto_status" not in sentence_summary.columns:
        return None
    valid = sentence_summary.dropna(subset=["majority_tag", "auto_status"])
    if valid.empty:
        return None
    return pd.crosstab(valid["majority_tag"], valid["auto_status"])


def compute_rater_profiles(resp_df: pd.DataFrame) -> pd.DataFrame:


    rows = []
    for rid, g in resp_df.groupby("rater_id"):
        tag_counts = g["level1_tag"].value_counts(normalize=True)
        row = {"rater_id": rid, "n_sentence_tags": g["level1_tag"].notna().sum()}
        for cat in LEVEL1_CATEGORIES:
            row[f"tag_{cat}_pct"] = round(tag_counts.get(cat, 0.0) * 100, 1)
        l2 = g.groupby(["item_id", "rater_id"], as_index=False).first()
        for dim in LEVEL2_DIMS:
            row[f"mean_{dim}"] = round(l2[dim].mean(), 2) if l2[dim].notna().any() else None
        rows.append(row)
    return pd.DataFrame(rows)


def compute_negative_control_detection(resp_df: pd.DataFrame, key_df: pd.DataFrame) -> dict:
    nc_ids = set(key_df[key_df["is_negative_control"]].index)
    nc_resp = resp_df[resp_df["item_id"].isin(nc_ids)]
    if nc_resp.empty:
        return {"error": "부정 대조 문항 응답이 없습니다."}

    per_response = []
    for (item_id, rater_id), g in nc_resp.groupby(["item_id", "rater_id"]):
        tags = g["level1_tag"].dropna()
        majority_unsupported = (not tags.empty) and (tags.mode().iloc[0] == "unsupported")
        faith = g["faithfulness"].dropna()
        low_faithfulness = (not faith.empty) and (faith.iloc[0] <= 2)
        detected = bool(majority_unsupported or low_faithfulness)
        per_response.append({"item_id": item_id, "rater_id": rater_id, "detected": detected})

    per_df = pd.DataFrame(per_response)
    detection_rate = per_df["detected"].mean()
    by_rater = per_df.groupby("rater_id")["detected"].mean()
    return {
        "overall_detection_rate": float(detection_rate),
        "n_control_responses": len(per_df),
        "by_rater": by_rater.to_dict(),
        "gate_passed": bool(detection_rate >= NEGATIVE_CONTROL_GATE),
    }


def find_discrepancy_candidates(item_summary: pd.DataFrame,
                                 auto_high=0.7, expert_low=2.0,
                                 auto_low=0.3, expert_high=4.0) -> pd.DataFrame:
    valid = item_summary.dropna(subset=["mean_faithfulness", "faithfulness_rate"])
    auto_missed = valid[
        (valid["faithfulness_rate"] >= auto_high) & (valid["mean_faithfulness"] <= expert_low)
    ].copy()
    auto_missed["discrepancy_type"] = "자동지표 미탐지(자동=충실, 전문가=불충실)"

    auto_overflagged = valid[
        (valid["faithfulness_rate"] <= auto_low) & (valid["mean_faithfulness"] >= expert_high)
    ].copy()
    auto_overflagged["discrepancy_type"] = "자동지표 과잉플래그(자동=불충실, 전문가=충실)"

    return pd.concat([auto_missed, auto_overflagged], ignore_index=True)


def find_sentence_discrepancy_candidates(sentence_summary: pd.DataFrame) -> pd.DataFrame:


    if sentence_summary is None or sentence_summary.empty or "auto_status" not in sentence_summary.columns:
        return pd.DataFrame()
    valid = sentence_summary.dropna(subset=["majority_tag", "auto_status"]).copy()
    if valid.empty:
        return pd.DataFrame()

    auto_pass_like = valid["auto_status"].isin(["PASS", "PASS_NO_SIGNAL"])
    auto_fail_content = valid["auto_status"].isin(["FAIL_NLI", "FAIL_INVALID_REF"])

    missed = valid[auto_pass_like & (valid["majority_tag"] == "unsupported")].copy()
    missed["discrepancy_type"] = "자동지표 미탐지(자동=PASS, 전문가=근거없음)"

    overflagged = valid[auto_fail_content & (valid["majority_tag"] == "supported")].copy()
    overflagged["discrepancy_type"] = "자동지표 과잉플래그(자동=FAIL, 전문가=지지됨)"

    return pd.concat([missed, overflagged], ignore_index=True)


# ════════════════════════════════════════════════════════════════

# ════════════════════════════════════════════════════════════════

def format_report(item_summary, sentence_summary, reliability, convergent, sentence_crosstab,
                   nc_detection, item_discrepancies, sentence_discrepancies, rater_profiles) -> str:
    lines = []
    lines.append("MetaKT-Verba 평정 결과 분석 리포트 (Part A5)")
    lines.append("=" * 60)
    lines.append(f"응답이 확인된 문항 수: {len(item_summary)}")
    lines.append(f"문항당 평균 평정자 수: {item_summary['n_raters'].mean():.2f}")
    lines.append("")

    lines.append("[0] 평정자별 태그/점수 사용 경향 (kappa/ICC가 낮을 때, '과제가 어려움' vs "
                  "'평정자마다 기준이 다름'을 구분하는 진단용)")
    lines.append(rater_profiles.to_string(index=False))
    lines.append("  ※ 특정 평정자의 tag_supported_pct/tag_unsupported_pct가 다른 평정자와 크게 다르거나, "
                  "mean_* 점수가 유독 높거나/낮으면 '기준 차이'일 가능성이 큽니다.")
    lines.append("")

    lines.append("[1] 평정자간 신뢰도 (core set 대상)")
    fk = reliability.get("level1_fleiss_kappa", {})
    if "error" in fk:
        lines.append(f"  - Level 1 (Fleiss' kappa, 명목형·모든 불일치 동일취급): {fk['error']}")
    else:
        lines.append(f"  - Level 1 (Fleiss' kappa, 명목형·모든 불일치 동일취급): κ = {fk['kappa']:.3f} "
                      f"(subject {fk['n_subjects']}개, 평정자 {fk['n_raters_per_subject']}명, "
                      f"불완전 응답으로 제외된 subject {fk.get('n_subjects_dropped_incomplete', 0)}개)")
    wk = reliability.get("level1_weighted_kappa", {})
    if wk.get("weighted_kappa") is None:
        lines.append("  - Level 1 (순서형 가중카파, 보조지표): 계산 불가")
    else:
        lines.append(f"  - Level 1 (순서형 가중카파, 보조지표 — 지지됨>부분지지>근거없음 순서 반영, "
                      f"판단불가는 최대거리 취급): κw = {wk['weighted_kappa']:.3f} "
                      f"(평정자쌍 {wk['n_rater_pairs']}쌍 평균)")
        if fk.get('kappa') is not None and wk['weighted_kappa'] - fk['kappa'] > 0.15:
            lines.append("    ⚠ 가중카파가 일반 카파보다 뚜렷이 높습니다 — 불일치 대부분이 "
                          "'한 칸 차이'(예: 지지됨↔부분지지)였을 가능성이 큽니다. 완전히 반대되는 "
                          "판단(지지됨↔근거없음)으로 갈린 경우는 상대적으로 적었다는 뜻입니다.")
    icc_map = reliability.get("level2_icc", {})
    for dim in LEVEL2_DIMS:
        r = icc_map.get(dim, {})
        if "error" in r:
            lines.append(f"  - Level 2 {dim} (ICC(2,1)): {r['error']}")
        else:
            lines.append(f"  - Level 2 {dim} (ICC(2,1)): {r['icc']:.3f} "
                          f"(문항 {r['n_targets']}개, 평정자 {r['k_raters']}명)")
    lines.append("")

    lines.append("[2-가] 자동-전문가 수렴타당도 (문항 단위)")
    for label, key in [("전문가 충실도 vs 자동 mean_entailment", "vs_mean_entailment"),
                        ("전문가 충실도 vs 자동 faithfulness_rate", "vs_auto_faithfulness_rate")]:
        c = convergent.get(key, {})
        if c.get("n", 0) < 3:
            lines.append(f"  - {label}: 유효 데이터 부족 (n={c.get('n', 0)})")
        else:
            lines.append(f"  - {label}: Pearson r = {c['pearson_r']:.3f} (p={c['pearson_p']:.3g}), "
                          f"Spearman ρ = {c['spearman_rho']:.3f} (p={c['spearman_p']:.3g}), n={c['n']}")
    lines.append("")

    lines.append("[2-나] 자동-전문가 수렴타당도 (문장 단위, 신규 — 전문가 다수결 태그 x 자동 status)")
    if sentence_crosstab is None:
        lines.append("  - 해당사항 없음 (--sentence-answer-key 미지정이거나 조인된 문장이 없음)")
    else:
        lines.append(sentence_crosstab.to_string().replace("\n", "\n  "))
    lines.append("")

    lines.append("[3] 부정 대조 문항(negative control) 탐지율 — QC 게이트")
    if "error" in nc_detection:
        lines.append(f"  - {nc_detection['error']}")
    else:
        gate_str = "통과" if nc_detection["gate_passed"] else f"⚠ 미달 (기준 {NEGATIVE_CONTROL_GATE:.0%})"
        lines.append(f"  - 전체 탐지율: {nc_detection['overall_detection_rate']:.1%} "
                      f"(응답 {nc_detection['n_control_responses']}건) — {gate_str}")
        for rid, rate in sorted(nc_detection["by_rater"].items()):
            lines.append(f"    · {rid}: {rate:.1%}")
    lines.append("")

    lines.append(f"[4-가] 자동-전문가 불일치 후보 — 문항 단위: {len(item_discrepancies)}건")
    for _, row in item_discrepancies.iterrows():
        lines.append(f"  - {row['item_id']} ({row['discrepancy_type']}): "
                      f"자동 faithfulness_rate={row['faithfulness_rate']:.2f}, "
                      f"전문가 평균={row['mean_faithfulness']:.2f}")
    lines.append("")
    lines.append(f"[4-나] 자동-전문가 불일치 후보 — 문장 단위(신규): {len(sentence_discrepancies)}건")
    if len(sentence_discrepancies):
        for _, row in sentence_discrepancies.iterrows():
            lines.append(f"  - {row['item_id']}/{row['sentence_id']} ({row['discrepancy_type']}): "
                         f"auto_status={row['auto_status']}, 전문가 다수결={row['majority_tag']} "
                         f"(분포: {row.get('tag_distribution', '')})")
    lines.append("")
    return "\n".join(lines)


# ════════════════════════════════════════════════════════════════
# main
# ════════════════════════════════════════════════════════════════

def run(responses_dir: str, answer_key_path: str, sentence_answer_key_path: str, output_dir: str,
        exclude_raters: list = None):
    os.makedirs(output_dir, exist_ok=True)
    resp_df = load_responses(responses_dir)
    if exclude_raters:
        before_n = resp_df["rater_id"].nunique()
        resp_df = resp_df[~resp_df["rater_id"].isin(exclude_raters)].copy()
        print(f"[민감도 분석] {exclude_raters} 제외 — 평정자 {before_n}명 -> "
              f"{resp_df['rater_id'].nunique()}명, {len(resp_df)}행 남음")
    key_df = load_answer_key(answer_key_path)
    sentence_key_df = load_sentence_answer_key(sentence_answer_key_path) if sentence_answer_key_path else None

    merged = resp_df.join(key_df, on="item_id", how="left")
    merged.to_csv(os.path.join(output_dir, "merged_responses.csv"), index=False)

    item_summary = build_item_level_summary(resp_df, key_df)
    item_summary.to_csv(os.path.join(output_dir, "item_level_summary.csv"), index=False)

    sentence_summary = build_sentence_level_summary(resp_df, sentence_key_df)
    if not sentence_summary.empty:
        sentence_summary.to_csv(os.path.join(output_dir, "sentence_level_summary.csv"), index=False)

    reliability = compute_core_set_reliability(resp_df, key_df)
    convergent = compute_convergent_validity(item_summary)
    sentence_crosstab = compute_sentence_level_crosstab(sentence_summary)
    nc_detection = compute_negative_control_detection(resp_df, key_df)
    item_discrepancies = find_discrepancy_candidates(item_summary)
    sentence_discrepancies = find_sentence_discrepancy_candidates(sentence_summary)
    rater_profiles = compute_rater_profiles(resp_df)
    rater_profiles.to_csv(os.path.join(output_dir, "rater_profiles.csv"), index=False)

    pd.concat([
        item_discrepancies.assign(level="item"),
        sentence_discrepancies.assign(level="sentence"),
    ], ignore_index=True).to_csv(os.path.join(output_dir, "discrepancy_candidates.csv"), index=False)

    report = format_report(item_summary, sentence_summary, reliability, convergent, sentence_crosstab,
                            nc_detection, item_discrepancies, sentence_discrepancies, rater_profiles)
    with open(os.path.join(output_dir, "report.txt"), "w", encoding="utf-8") as f:
        f.write(report)
    print("\n" + report)
    print(f"\n✅ 분석 완료: {output_dir}")


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--responses-dir", type=str,
                         default='data/ratings',
                         help="평정자 회신 CSV/XLSX 파일들이 있는 폴더")
    parser.add_argument("--answer-key", type=str,
                         default='data/stimuli/internal_answer_key.csv',
                         help="internal_answer_key.csv 경로 (연구팀 전용, 평정자에게 공유 금지)")
    parser.add_argument("--sentence-answer-key", type=str,
                         default='data/stimuli/internal_answer_key_sentences.csv',
                         help="internal_answer_key_sentences.csv 경로 (연구팀 전용, 선택 — 없으면 문장 단위 "
                              "자동-전문가 비교[2-나/4-나]는 건너뜀)")
    parser.add_argument("--output-dir", type=str,
                         default="outputs/ratings_analysis")
    parser.add_argument("--exclude-raters", nargs="+", default=None, #"R1".split(),
                         help="민감도 분석용 — 지정한 rater_id를 전부 제외하고 재계산. "
                              "예: --exclude-raters R1 R3 (특정 평정자가 신뢰도 전체를 "
                              "끌어내리는지 확인할 때 사용, output-dir 아래 별도 하위폴더에 저장)")

    args = parser.parse_args()


    if not args.responses_dir or not args.answer_key:
        parser.error("--responses-dir와 --answer-key는 필수입니다.")

    run(args.responses_dir, args.answer_key, args.sentence_answer_key, args.output_dir)

    if args.exclude_raters:
        excl_tag = "_".join(args.exclude_raters)
        sensitivity_dir = os.path.join(args.output_dir, f"sensitivity_excl_{excl_tag}")
        print(f"\n{'=' * 60}\n민감도 분석: {args.exclude_raters} 제외\n{'=' * 60}")
        run(args.responses_dir, args.answer_key, args.sentence_answer_key, sensitivity_dir,
            exclude_raters=args.exclude_raters)
        print(f"\n※ 위 [전체]와 방금 [{excl_tag} 제외] 결과의 [1]번(신뢰도) 수치를 비교해보세요 — "
              f"제외 후 kappa/ICC가 뚜렷이 오르면 해당 평정자들이 신뢰도를 끌어내린 것이고, "
              f"큰 차이가 없으면 낮은 신뢰도가 과제 자체의 어려움일 가능성이 큽니다.")


if __name__ == "__main__":
    main()